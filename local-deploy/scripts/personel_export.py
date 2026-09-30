# -*- coding: utf-8 -*-
"""
PERSONEL LISTESI (ISE GIRIS / ISTEN CIKIS) DISA AKTARIMI
=========================================================
server.py ile AYNI container icinde calisir (daily_leave_log_export.py gibi).
Personel Canli Takip uygulamasinin "Devam Takvimi" sayfasi, puantaji dogru
hesaplamak icin personelin gercek ise giris ve isten cikis tarihlerini bu
dosyadan okur.

URETILEN DOSYA (container icinde, bat script host'a kopyalar):
  /tmp/Personel_Listesi.xlsx
  Kolonlar: Sicil No | Ad Soyad | Ise Giris | Isten Cikis | Aktif | Departman
  Tarihler YYYY-MM-DD formatindadir.

Veritabanina hicbir sey YAZMAZ, sadece okur.
"""
import asyncio
import sys
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, "/app")
from server import db  # noqa: E402

OUT_PATH = "/tmp/Personel_Listesi.xlsx"


def iso(v):
    if not v:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    for fmt in ("%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


async def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Personel"
    header = ["Sicil No", "Ad Soyad", "Ise Giris", "Isten Cikis", "Aktif", "Departman"]
    ws.append(header)
    for i, w in enumerate((12, 28, 13, 13, 8, 26), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = w
    n = 0
    async for p in db.personnel.find({}, {"_id": 0, "sicil_no": 1, "ad_soyad": 1, "ise_giris": 1,
                                         "isten_cikis": 1, "aktif": 1, "departman": 1}):
        ws.append([
            str(p.get("sicil_no") or ""),
            p.get("ad_soyad") or "",
            iso(p.get("ise_giris")),
            iso(p.get("isten_cikis")),
            "Evet" if p.get("aktif", True) else "Hayir",
            p.get("departman") or "",
        ])
        n += 1
    ws.freeze_panes = "A2"
    wb.save(OUT_PATH)
    print(f"OK: Personel listesi -> {n} kisi yazildi.")


if __name__ == "__main__":
    asyncio.run(main())
